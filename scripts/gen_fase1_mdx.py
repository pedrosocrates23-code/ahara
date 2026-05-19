import sys, re, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

OUTPUT_DIR = r"D:\ia\Projeto Ahara\src\content\blog"
PUB_DATE   = "2026-05-19T10:00:00.000Z"

def strip_tags(html):
    return re.sub(r'<[^>]+>', '', html or '').strip()

def extract_faqs(html):
    faqs = []
    faq_section = re.search(
        r'<h2[^>]*>Perguntas frequentes[^<]*</h2>(.*?)(?:<h2|$)',
        html, re.DOTALL | re.IGNORECASE
    )
    block = faq_section.group(1) if faq_section else None
    if not block:
        all_uls = list(re.finditer(r'<ul[^>]*>(.*?)</ul>', html, re.DOTALL | re.IGNORECASE))
        block = all_uls[-1].group(1) if all_uls else None
    if not block:
        return faqs
    for item in re.findall(r'<li[^>]*>(.*?)</li>', block, re.DOTALL):
        q_match = re.search(r'<strong[^>]*>(.*?)</strong>', item)
        if not q_match:
            continue
        question = strip_tags(q_match.group(1)).strip()
        if not question.endswith('?'):
            question += '?'
        answer_html = re.sub(r'^.*?</strong>\s*', '', item, count=1, flags=re.DOTALL)
        answer = ' '.join(strip_tags(answer_html).split())
        question = question.replace('"', "'")
        answer   = answer.replace('"', "'")
        if question and answer:
            faqs.append({"q": question, "a": answer})
    return faqs

def ys(s):
    escaped = s.replace('"', "'")
    return '"' + escaped + '"'

def clean_seo_title(s):
    return s.strip().rstrip("\\").strip()

def build_mdx(slug, h1, seo_title, keyword, meta_desc, sumario_html, palavras, faqs, html_body):
    lines = ['---']
    lines.append('h1: ' + ys(h1))
    seo_clean = clean_seo_title(seo_title)
    if seo_clean and seo_clean != h1:
        lines.append('seo_title: ' + ys(seo_clean))
    lines.append('keyword: ' + ys(keyword))
    lines.append('meta_description: ' + ys(meta_desc))
    if sumario_html:
        lines.append('sumario_html: ' + ys(sumario_html))
    lines.append(f'palavras_totais: {palavras}')
    lines.append('pub_date: "' + PUB_DATE + '"')
    if faqs:
        lines.append('faq:')
        for f in faqs:
            lines.append('  - q: ' + ys(f['q']))
            lines.append('    a: ' + ys(f['a']))
    lines.append('---')
    lines.append('')
    lines.append(html_body.strip())
    return '\n'.join(lines)

wb = openpyxl.load_workbook(r"C:\Users\User\Downloads\ahara-fase1-artigos-cms.xlsx")
ws = wb['Ahara Fase 1']

criados, erros = [], []

for r in range(2, ws.max_row + 1):
    slug      = (ws.cell(r,  1).value or '').strip()
    keyword   = (ws.cell(r,  2).value or '').strip()
    h1        = (ws.cell(r,  5).value or '').strip()
    seo_title = (ws.cell(r,  6).value or '').strip()
    meta_desc = (ws.cell(r,  7).value or '').strip()
    palavras  = ws.cell(r,  9).value or ''
    sumario   = (ws.cell(r, 10).value or '').strip()
    html_body = (ws.cell(r, 11).value or '').strip()
    if not slug:
        continue

    wc_match = re.search(r'[\d.,]+', str(palavras))
    wc = int(re.sub(r'[.,]', '', wc_match.group())) if wc_match else 1200

    out_path = os.path.join(OUTPUT_DIR, slug + '.mdx')
    if os.path.exists(out_path):
        erros.append(f"SKIP (ja existe): {slug}")
        continue

    faqs = extract_faqs(html_body)
    mdx  = build_mdx(slug, h1, seo_title, keyword, meta_desc, sumario, wc, faqs, html_body)

    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(mdx)
    criados.append((slug, wc, len(faqs)))

print(f"Criados: {len(criados)} | Pulados: {len(erros)}")
print(f"\n{'#':<3} {'Slug':<46} {'Words':<7} {'FAQs'}")
print("-" * 65)
for i, (slug, wc, fc) in enumerate(criados, 1):
    print(f"{i:<3} {slug:<46} {wc:<7} {fc}")
if erros:
    print("\nPulados:")
    for e in erros:
        print(" ", e)
